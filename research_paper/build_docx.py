"""Generate an editable .docx version of the paper from paper_content.

Single-column (editable) layout: headings, bold abstract, tables, embedded
figures and equation images, and a numbered reference list. Not a two-column
typeset reproduction - it is the editable Word deliverable.
Run:  python research_paper/build_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
sys.path.insert(0, str(HERE))
import paper_content as C  # noqa: E402

from PIL import Image as PILImage


def add_rich(p, text):
    """Add text to a paragraph, honoring *italic* markup; [n] kept literal."""
    for i, seg in enumerate(re.split(r"\*(.+?)\*", text)):
        run = p.add_run(seg)
        if i % 2 == 1:
            run.italic = True


def pic_width(img: Path, wide: bool) -> Inches:
    iw, ih = PILImage.open(img).size
    target = 6.3 if wide else 3.1
    return Inches(min(target, iw / 300.0))


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(C.TITLE); r.bold = True; r.font.size = Pt(16)

    a = C.AUTHORS[0]
    for line in (a["name"], a["dept"], a["org"], a["extra"]):
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.add_run(line)

    pa = doc.add_paragraph()
    pa.add_run("Abstract—").bold = True
    add_rich(pa, C.ABSTRACT)
    pk = doc.add_paragraph()
    pk.add_run("Index Terms—").bold = True
    pk.add_run(C.INDEX)

    for b in C.BLOCKS:
        k = b[0]
        if k == "h1":
            doc.add_heading(f"{b[1]}. {b[2].title()}", level=1)
        elif k == "h2":
            doc.add_heading(f"{b[1]}. {b[2]}", level=2)
        elif k == "p":
            add_rich(doc.add_paragraph(), b[1])
        elif k == "runin":
            p = doc.add_paragraph(); p.add_run(b[1] + " ").italic = True
            add_rich(p, b[2])
        elif k == "olist":
            for it in b[1]:
                add_rich(doc.add_paragraph(style="List Number"), it)
        elif k == "ulist":
            for it in b[1]:
                add_rich(doc.add_paragraph(style="List Bullet"), it)
        elif k == "eq":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(FIG / f"eq{b[1]}.png"), width=Inches(2.6))
            p.add_run(f"    ({b[1]})")
        elif k == "fig":
            _, label, cap, img, wide = b
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(FIG / img), width=pic_width(FIG / img, wide))
            cp = doc.add_paragraph()
            cp.add_run(f"Fig. {label}. ").bold = True
            cp.add_run(cap)
        elif k == "table":
            _, label, cap, header, rows, wide = b
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.add_run(f"TABLE {label}. {cap}").bold = True
            tb = doc.add_table(rows=1, cols=len(header))
            tb.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                tb.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
            for row in rows:
                cells = tb.add_row().cells
                for j, c in enumerate(row):
                    cells[j].text = str(c)
        elif k == "codeavail":
            doc.add_heading("Code Availability", level=1)
            add_rich(doc.add_paragraph(), b[1])

    doc.add_heading("References", level=1)
    for i, ref in enumerate(C.REFERENCES, 1):
        doc.add_paragraph(f"[{i}] {ref}")

    doc.save(str(HERE / "birdvision_research_paper.docx"))
    print("Wrote birdvision_research_paper.docx")


if __name__ == "__main__":
    main()
